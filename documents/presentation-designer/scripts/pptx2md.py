"""Convert PPTX to Markdown presentation.

Usage:
    python pptx2md.py input.pptx [output.md]

Extracts slide titles, content text, and speaker notes, then formats them
into the skill's Markdown format with ASCII art boxes.

Requires: python-pptx (pip install python-pptx)
"""

import re
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.enum.text import PP_ALIGN
except ImportError:
    print("Error: python-pptx is required. Install with: pip install python-pptx")
    sys.exit(1)


BOX_WIDTH = 66
SLIDE_TYPE_PREFIX = "slide-type:"


def _marker(alignment):
    """Map a PPTX paragraph alignment to the skill's inline marker.

    LEFT (and inherited) yields no marker because LEFT is the default;
    RIGHT -> `>>`, CENTER -> `=`.
    """
    if alignment == PP_ALIGN.RIGHT:
        return ">>"
    if alignment == PP_ALIGN.CENTER:
        return "="
    return ""


def detect_slide_type(slide):
    """Recover the slide type stored by `md2pptx.py` (returns None if absent)."""
    for shape in slide.shapes:
        name = getattr(shape, "name", "") or ""
        if name.startswith(SLIDE_TYPE_PREFIX):
            return name[len(SLIDE_TYPE_PREFIX):].strip() or None
    return None


def extract_slide_lines(slide):
    """Return [(text, pptx_alignment), ...] for each non-empty paragraph."""
    out = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                t = para.text.strip()
                if t:
                    out.append((t, para.alignment))
    return out


def get_speaker_notes(slide):
    try:
        notes_slide = slide.notes_slide
        return notes_slide.notes_text_frame.text.strip()
    except Exception:
        return ""


def guess_title(texts):
    if not texts:
        return "Untitled"
    # First non-empty line that looks like a title (short, no leading dash/bullet)
    for t in texts:
        if len(t) < 120 and not t.startswith("-") and not t.startswith("•"):
            return t
    return texts[0]


def make_ascii_box(lines):
    top = "+" + "-" * (BOX_WIDTH - 2) + "+"
    bottom = top
    result = [top]
    
    # Ensure exactly 13 lines of content (15 total - 2 borders)
    # In a real implementation, we'd split the content into header and body
    # For now, we pad with empty lines to reach the fixed height
    for i in range(13):
        if i < len(lines):
            line = lines[i][:BOX_WIDTH - 4]
            result.append("| " + line.ljust(BOX_WIDTH - 4) + " |")
        else:
            result.append("| " + " " * (BOX_WIDTH - 4) + " |")
            
    result.append(bottom)
    return "\n".join(result)


def wrap_text(text, width=BOX_WIDTH - 4):
    words = text.split()
    lines = []
    current = ""
    for w in words:
        if len(current) + len(w) + 1 <= width:
            current = (current + " " + w).strip()
        else:
            if current:
                lines.append(current)
            current = w
    if current:
        lines.append(current)
    return lines if lines else [text]


def wrap_with_marker(text, alignment):
    """Wrap a paragraph, prefixing every line with its alignment marker."""
    marker = _marker(alignment)
    wrapped = wrap_text(text, width=BOX_WIDTH - 4 - len(marker))
    return [marker + line for line in wrapped]


def convert(pptx_path, md_path):
    prs = Presentation(pptx_path)
    slides = prs.slides
    md_lines = []
    slide_num = 0

    for slide in slides:
        lines_info = extract_slide_lines(slide)
        if not lines_info:
            continue
        notes = get_speaker_notes(slide)

        slide_type = detect_slide_type(slide)
        texts = [t for t, _ in lines_info]
        title = guess_title(texts)

        title_index = next(i for i, (t, _) in enumerate(lines_info) if t == title)
        title_text, title_alignment = lines_info[title_index]
        lines = [_marker(title_alignment) + title_text]
        for i, (t, al) in enumerate(lines_info):
            if i != title_index:
                lines.extend(wrap_with_marker(t, al))

        slide_num += 1
        md_lines.append(f"## Slide {slide_num}. {title}")
        md_lines.append("")
        fence = f"```slide:{slide_type}" if slide_type else "```slide"
        md_lines.append(fence)
        md_lines.append(make_ascii_box(lines))
        md_lines.append("```")
        md_lines.append("")
        if notes:
            md_lines.append(f"__Comment:__ {notes}")
        else:
            md_lines.append("__Comment:__")
        md_lines.append("")

    content = "\n".join(md_lines)
    md_path.write_text(content, encoding="utf-8")
    print(f"Converted {slide_num} slides -> {md_path}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    pptx_path = Path(sys.argv[1])
    if not pptx_path.exists():
        print(f"File not found: {pptx_path}")
        sys.exit(1)

    md_path = Path(sys.argv[2]) if len(sys.argv) > 2 else pptx_path.with_suffix(".md")
    convert(pptx_path, md_path)
