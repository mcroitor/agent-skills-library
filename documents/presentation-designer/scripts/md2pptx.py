"""Convert Markdown presentation to PPTX.

Usage:
    python md2pptx.py input.md [output.pptx] [--style NAME]

Parses slides in the format:
    ## Slide N. Title
    ```slide[:type]
    +-- ASCII art box --+
    | content           |
    +-------------------+
    ```
    __Comment:__ Speaker notes (may span multiple lines)

Slide types: title, section, content (default), two-columns.
Styles: academic (default), dark, minimal — see styles.json.

Requires: python-pptx (pip install python-pptx)
"""

import json
import re
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
except ImportError:
    print("Error: python-pptx is required. Install with: pip install python-pptx")
    sys.exit(1)

STYLES_PATH = Path(__file__).resolve().parent.parent / "assets/styles.json"
DEFAULT_STYLE = "academic"
SLIDE_TYPE_PREFIX = "slide-type:"


def load_styles():
    with open(STYLES_PATH, encoding="utf-8") as f:
        return json.load(f)


def parse_markdown(filepath):
    """Parse the Markdown file into a list of slide dicts.

    Each slide dict has: number, title, type, slide_content, comment.
    The comment captures ALL lines after `__Comment:__` up to the next
    `## Slide` header or end of file, preserving internal newlines.

    Implemented as a chunk-based parser (split on slide headers) rather than
    one large regex, which is far more robust against whitespace variations.
    """
    text = Path(filepath).read_text(encoding="utf-8-sig")

    # Split into chunks, each starting at a `## Slide N. Title` header.
    header_re = re.compile(r"^##\s+Slide\s+(\d+)\.\s+(.+?)\s*$", re.MULTILINE)
    headers = list(header_re.finditer(text))
    if not headers:
        return []

    slides = []
    for i, h in enumerate(headers):
        number = h.group(1)
        title = h.group(2).strip()
        start = h.end()
        end = headers[i + 1].start() if i + 1 < len(headers) else len(text)
        chunk = text[start:end]

        # Extract the fenced slide block and its type.
        fence_re = re.compile(
            r"^```slide(?::([\w-]+))?\s*\n(.*?)^```\s*$",
            re.MULTILINE | re.DOTALL,
        )
        fence = fence_re.search(chunk)
        if fence:
            slide_type = fence.group(1) or "content"
            content = fence.group(2).strip()
        else:
            slide_type = "content"
            content = ""

        line_count = len(content.splitlines())
        if line_count != 15:
            print(
                f"Warning: Slide {number} has {line_count} "
                f"lines instead of 15."
            )

        # Extract the comment: everything after the `__Comment:__` marker
        # up to the next slide header or end of file, then strip the leading
        # underscores/whitespace and any trailing blank lines.
        comment = ""
        cidx = chunk.find("__Comment")
        if cidx != -1:
            colon = chunk.find(":", cidx)
            if colon != -1:
                comment = re.sub(r"^[\s_]+", "", chunk[colon + 1:]).strip()

        slides.append(
            {
                "number": number,
                "title": title,
                "type": slide_type,
                "slide_content": content,
                "comment": comment,
            }
        )
    return slides


def _clean_line(line):
    """Strip box borders and alignment markers from one ASCII line.

    Returns (text, alignment) or None if the line is empty/border-only.
    Border pipes (`|`) are stripped independently from both ends and
    repeatedly, so nested boxes (e.g. the section template's inner box)
    collapse down to their innermost content.
    Alignment markers: `<<Text` -> LEFT, `>>Text` -> RIGHT, `=Text` -> CENTER 
        default is LEFT if no marker is present.
    """
    s = line.strip()
    if not s:
        return None
    if set(s) <= set("+-|"):
        return None  # pure border line

    # Strip border pipes from both ends (independently), handling nesting.
    changed = True
    while changed:
        changed = False
        if s.startswith("|"):
            s = s[1:].strip()
            changed = True
        if s.endswith("|"):
            s = s[:-1].strip()
            changed = True
        if not s or set(s) <= set("+-|"):
            return None

    alignment = PP_ALIGN.LEFT  # default is LEFT if no marker is present
    if s.startswith("<<"):
        alignment = PP_ALIGN.LEFT
        s = s[2:].strip()
    elif s.startswith(">>"):
        alignment = PP_ALIGN.RIGHT
        s = s[2:].strip()
    elif s.startswith("="):
        alignment = PP_ALIGN.CENTER
        s = s[1:].strip()
    return s, alignment


def extract_text_lines(ascii_content):
    """Extract non-empty text lines (with alignment) from a single-column box."""
    result = []
    for line in ascii_content.splitlines():
        cleaned = _clean_line(line)
        if cleaned:
            result.append(cleaned)
    return result


def extract_two_columns(ascii_content):
    """Split a two-column box into (header, (left_lines, right_lines)).

    Detects the `+---+---+` split line; text above it is the shared header,
    text below is split on the middle `|`.
    """
    lines = ascii_content.splitlines()
    split_idx = None
    for i, line in enumerate(lines):
        s = line.strip()
        if s.startswith("+") and s.endswith("+") and s.count("+") >= 3:
            split_idx = i
            break
    if split_idx is None:
        # No split line found — fall back to single-column extraction.
        return [], (extract_text_lines(ascii_content), [])

    header = extract_text_lines("\n".join(lines[:split_idx]))
    left, right = [], []
    for line in lines[split_idx + 1:]:
        s = line.strip()
        if s.startswith("+") and s.endswith("+"):
            continue
        if "|" not in s:
            continue
        # The middle pipe is the one closest to the horizontal center.
        mid = len(s) // 2
        pipes = [i for i, ch in enumerate(s) if ch == "|"]
        if len(pipes) < 2:
            continue
        best = min(pipes, key=lambda i: abs(i - mid))
        left_raw = s[:best]
        right_raw = s[best + 1:]
        lc = _clean_line(left_raw)
        rc = _clean_line(right_raw)
        if lc:
            left.append(lc)
        if rc:
            right.append(rc)
    return header, (left, right)


def _style_run(run, style, size, bold=False, italic=False):
    run.font.name = style["font"]
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(style["foreground"].lstrip("#"))


def _add_paragraphs(tf, lines, style, size, bold=False, italic=False, space_after=6):
    for i, (text, alignment) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.alignment = alignment
        p.space_after = Pt(space_after)
        for run in p.runs:
            _style_run(run, style, size, bold=bold, italic=italic)


def _apply_background(slide, style):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(style["background"].lstrip("#"))


def _tag_slide_type(slide, slide_type):
    """Store the slide type on a shape name so `pptx2md.py` can recover it.

    The type is not part of the rendered slide; it is attached to the first
    shape's name (or an off-canvas marker shape when the slide is empty).
    """
    shapes = list(slide.shapes)
    if shapes:
        shapes[0].name = f"{SLIDE_TYPE_PREFIX}{slide_type}"
    else:
        marker = slide.shapes.add_textbox(
            Inches(-100), Inches(-100), Inches(1), Inches(1)
        )
        marker.name = f"{SLIDE_TYPE_PREFIX}{slide_type}"


def add_slide(prs, slide_data, style):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    _apply_background(slide, style)

    stype = slide_data["type"]
    content = slide_data["slide_content"]

    if stype == "title":
        lines = extract_text_lines(content)
        if lines:
            main = lines[0]
            rest = lines[1:]
            box = slide.shapes.add_textbox(Inches(0.5), Inches(1.8), Inches(9), Inches(1.5))
            tf = box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = main[0]
            p.alignment = main[1]
            for run in p.runs:
                _style_run(run, style, 32, bold=True)
            if rest:
                sub = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9), Inches(2.0))
                stf = sub.text_frame
                stf.word_wrap = True
                _add_paragraphs(stf, rest, style, 16)

    elif stype == "section":
        lines = extract_text_lines(content)
        if lines:
            main = lines[0]
            rest = lines[1:]
            box = slide.shapes.add_textbox(Inches(0.5), Inches(2.2), Inches(9), Inches(1.5))
            tf = box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = main[0]
            p.alignment = main[1]
            for run in p.runs:
                _style_run(run, style, 30, bold=True)
            if rest:
                moto = rest[-1]
                m = slide.shapes.add_textbox(Inches(0.5), Inches(4.5), Inches(9), Inches(1.0))
                mtf = m.text_frame
                mtf.word_wrap = True
                p2 = mtf.paragraphs[0]
                p2.text = moto[0]
                p2.alignment = moto[1]
                for run in p2.runs:
                    _style_run(run, style, 16, italic=True)

    elif stype == "two-columns":
        header, (left, right) = extract_two_columns(content)
        if header:
            h = header[0]
            hbox = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.7))
            htf = hbox.text_frame
            htf.word_wrap = True
            p = htf.paragraphs[0]
            p.text = h[0]
            p.alignment = h[1]
            for run in p.runs:
                _style_run(run, style, 24, bold=True)
        if left:
            lb = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(4.3), Inches(5.5))
            ltf = lb.text_frame
            ltf.word_wrap = True
            _add_paragraphs(ltf, left, style, 16)
        if right:
            rb = slide.shapes.add_textbox(Inches(5.2), Inches(1.3), Inches(4.3), Inches(5.5))
            rtf = rb.text_frame
            rtf.word_wrap = True
            _add_paragraphs(rtf, right, style, 16)

    else:  # content (default)
        lines = extract_text_lines(content)
        if lines:
            header = lines[0]
            body = lines[1:]
            hbox = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.7))
            htf = hbox.text_frame
            htf.word_wrap = True
            p = htf.paragraphs[0]
            p.text = header[0]
            p.alignment = header[1]
            for run in p.runs:
                _style_run(run, style, 24, bold=True)
            if body:
                cbox = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(9), Inches(5.5))
                ctf = cbox.text_frame
                ctf.word_wrap = True
                _add_paragraphs(ctf, body, style, 18)

    _tag_slide_type(slide, stype)

    # --- Speaker notes (multi-line preserved) ---
    comment = slide_data["comment"]
    if comment:
        notes_tf = slide.notes_slide.notes_text_frame
        notes_tf.text = comment


def convert(md_path, pptx_path, style_name):
    styles = load_styles()
    if style_name not in styles:
        available = ", ".join(styles.keys())
        print(f"Error: unknown style '{style_name}'. Available: {available}")
        sys.exit(1)
    style = styles[style_name]

    slides = parse_markdown(md_path)
    if not slides:
        print(f"No slides found in {md_path}")
        sys.exit(1)

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    for s in slides:
        add_slide(prs, s, style)

    prs.save(pptx_path)
    print(f"Converted {len(slides)} slides -> {pptx_path} (style: {style_name})")


if __name__ == "__main__":
    args = sys.argv[1:]
    style_name = DEFAULT_STYLE
    if "--style" in args:
        idx = args.index("--style")
        if idx + 1 >= len(args):
            print("Error: --style requires a value.")
            sys.exit(1)
        style_name = args[idx + 1]
        del args[idx: idx + 2]

    if not args:
        print(__doc__)
        sys.exit(1)

    md_path = Path(args[0])
    if not md_path.exists():
        print(f"File not found: {md_path}")
        sys.exit(1)

    pptx_path = Path(args[1]) if len(args) > 1 else md_path.with_suffix(".pptx")
    convert(md_path, pptx_path, style_name)
